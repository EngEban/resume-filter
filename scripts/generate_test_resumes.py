# ============================================================
# scripts/generate_test_resumes.py
# Generate 12 diverse test resumes for B2B batch testing.
# ============================================================
import os
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH


OUTPUT_DIR = Path("/tmp/test_resumes")
OUTPUT_DIR.mkdir(exist_ok=True, parents=True)


# ============================================================
# 12 Candidate Profiles (varied match quality)
# ============================================================
CANDIDATES = [
    {
        "name": "Ahmad Khalil",
        "email": "ahmad.khalil@example.com",
        "phone": "+970 599 111 111",
        "location": "Gaza",
        "summary": "Senior Python Backend Developer with 8 years of experience.",
        "experience": [
            "Lead Backend Engineer at TechCorp (2020-Present)",
            "- Architected microservices with FastAPI handling 20M requests/day",
            "- Optimized PostgreSQL queries reducing latency by 70%",
            "- Led team of 6 engineers across 3 time zones",
            "- Deployed Kubernetes clusters on AWS and GCP",
            "Senior Developer at StartupX (2017-2020)",
            "- Built REST APIs with Django REST Framework",
            "- Implemented Celery-based task queues processing 1M jobs/day",
        ],
        "education": "BSc in Computer Science, Islamic University of Gaza (2016)",
        "skills": "Python, FastAPI, PostgreSQL, Redis, Docker, Kubernetes, Celery, AWS, REST APIs, GraphQL, Git, CI/CD",
        "match": "excellent",
    },
    {
        "name": "Sara Mansour",
        "email": "sara.mansour@example.com",
        "phone": "+970 599 222 222",
        "location": "Ramallah",
        "summary": "Backend Developer with 5 years of Python experience.",
        "experience": [
            "Backend Developer at FinTechCo (2021-Present)",
            "- Developed microservices with FastAPI and PostgreSQL",
            "- Built Docker containerized services",
            "- Implemented Redis caching layer reducing DB load by 40%",
            "Junior Developer at WebAgency (2019-2021)",
            "- Wrote Python scripts for data processing",
        ],
        "education": "BSc in Software Engineering, Birzeit University (2018)",
        "skills": "Python, FastAPI, PostgreSQL, Redis, Docker, Git",
        "match": "good",
    },
    {
        "name": "Omar Yousef",
        "email": "omar.yousef@example.com",
        "phone": "+970 599 333 333",
        "location": "Nablus",
        "summary": "Full-Stack Developer with 4 years experience.",
        "experience": [
            "Full-Stack Developer at E-commerce Co (2020-Present)",
            "- Built Node.js and Express APIs",
            "- Worked with MongoDB and MySQL",
            "- Frontend work with React",
        ],
        "education": "BSc in Computer Science, An-Najah University (2019)",
        "skills": "JavaScript, Node.js, React, MongoDB, MySQL, Git",
        "match": "weak",
    },
    {
        "name": "Layla Hasan",
        "email": "layla.hasan@example.com",
        "phone": "+970 599 444 444",
        "location": "Hebron",
        "summary": "Python Developer with 6 years of backend experience.",
        "experience": [
            "Backend Engineer at CloudServices (2018-Present)",
            "- Built scalable APIs with FastAPI",
            "- Managed PostgreSQL databases with high availability",
            "- Implemented Celery for background processing",
            "- Used Redis for caching and task queues",
        ],
        "education": "MSc in Computer Science, Hebron University (2017)",
        "skills": "Python, FastAPI, PostgreSQL, Redis, Celery, Docker, Kubernetes, Git, Linux",
        "match": "excellent",
    },
    {
        "name": "Youssef Ali",
        "email": "youssef.ali@example.com",
        "phone": "+970 599 555 555",
        "location": "Gaza",
        "summary": "Backend Developer with 3 years experience.",
        "experience": [
            "Developer at SoftwareHouse (2021-Present)",
            "- Worked with Python Flask",
            "- Basic PostgreSQL experience",
        ],
        "education": "Diploma in IT, UCAS (2020)",
        "skills": "Python, Flask, PostgreSQL, Git",
        "match": "below_average",
    },
    {
        "name": "Nour Ibrahim",
        "email": "nour.ibrahim@example.com",
        "phone": "+970 599 666 666",
        "location": "Bethlehem",
        "summary": "Senior Software Engineer with 10 years experience.",
        "experience": [
            "Principal Engineer at BigTech (2015-Present)",
            "- Designed distributed systems with Python",
            "- Expert in FastAPI, PostgreSQL, Kubernetes",
            "- Mentored 15+ engineers",
            "- Published papers on scalable architecture",
            "- Speaker at PyCon MEA 2022",
        ],
        "education": "PhD in Computer Science, Technion (2014)",
        "skills": "Python, FastAPI, PostgreSQL, Kubernetes, Docker, Redis, Celery, AWS, GCP, Terraform, gRPC, Kafka",
        "match": "excellent",
    },
    {
        "name": "Rami Saleh",
        "email": "rami.saleh@example.com",
        "phone": "+970 599 777 777",
        "location": "Jenin",
        "summary": "Junior Developer looking for opportunities.",
        "experience": [
            "Intern at Startup (2023-2024)",
            "- Learned Python basics",
            "- Assisted with bug fixes",
        ],
        "education": "BSc in IT, Arab American University (2023)",
        "skills": "Python, HTML, CSS, Git",
        "match": "weak",
    },
    {
        "name": "Dina Khaled",
        "email": "dina.khaled@example.com",
        "phone": "+970 599 888 888",
        "location": "Ramallah",
        "summary": "Backend Engineer specializing in API development.",
        "experience": [
            "Backend Engineer at SaaS Company (2019-Present)",
            "- Developed RESTful APIs with FastAPI",
            "- Managed PostgreSQL and Redis infrastructure",
            "- Set up Docker and CI/CD pipelines",
            "- Worked with Celery for async tasks",
            "- Deployed on Kubernetes",
        ],
        "education": "BSc in Computer Engineering, Birzeit University (2018)",
        "skills": "Python, FastAPI, PostgreSQL, Redis, Docker, Kubernetes, Celery, Git, Jenkins",
        "match": "good",
    },
    {
        "name": "Khaled Nasser",
        "email": "khaled.nasser@example.com",
        "phone": "+970 599 999 999",
        "location": "Gaza",
        "summary": "DevOps Engineer with Python background.",
        "experience": [
            "DevOps Engineer at CloudInfra (2018-Present)",
            "- Managed Kubernetes clusters",
            "- Automated deployments with Terraform",
            "- Wrote Python automation scripts",
        ],
        "education": "BSc in Computer Science, IUG (2017)",
        "skills": "Python, Kubernetes, Docker, Terraform, AWS, Linux, Git, CI/CD",
        "match": "average",
    },
    {
        "name": "Huda Sami",
        "email": "huda.sami@example.com",
        "phone": "+970 599 000 000",
        "location": "Nablus",
        "summary": "Python Developer with strong data background.",
        "experience": [
            "Data Engineer at Analytics Co (2020-Present)",
            "- Built data pipelines with Python",
            "- Worked with PostgreSQL and Airflow",
            "- Implemented Redis caching",
        ],
        "education": "BSc in Statistics, An-Najah University (2019)",
        "skills": "Python, PostgreSQL, Redis, Airflow, Pandas, NumPy, SQL",
        "match": "average",
    },
    {
        "name": "Tariq Zaid",
        "email": "tariq.zaid@example.com",
        "phone": "+970 599 121 212",
        "location": "Hebron",
        "summary": "Senior Backend Developer with 7 years of Python experience.",
        "experience": [
            "Senior Developer at Enterprise (2017-Present)",
            "- Built FastAPI microservices",
            "- Expert in PostgreSQL performance tuning",
            "- Implemented Redis and Celery workflows",
            "- Led Kubernetes migration",
            "- Mentored 4 junior developers",
            "- 99.99% uptime on critical services",
        ],
        "education": "BSc in Software Engineering, Polytechnic University (2016)",
        "skills": "Python, FastAPI, PostgreSQL, Redis, Celery, Docker, Kubernetes, Git, Prometheus, Grafana",
        "match": "excellent",
    },
    {
        "name": "Maya Adel",
        "email": "maya.adel@example.com",
        "phone": "+970 599 343 434",
        "location": "Ramallah",
        "summary": "Software Developer with 2 years experience.",
        "experience": [
            "Developer at Small Agency (2022-Present)",
            "- Built web apps with Python and Flask",
            "- Some PostgreSQL experience",
            "- Used Docker occasionally",
        ],
        "education": "BSc in CS, Al-Quds University (2022)",
        "skills": "Python, Flask, PostgreSQL, Docker, Git",
        "match": "below_average",
    },
]


# ============================================================
# DOCX Generator
# ============================================================
def create_resume_docx(candidate: dict, output_path: Path) -> None:
    doc = Document()

    # Name
    name_para = doc.add_heading(candidate["name"], level=0)
    name_para.alignment = WD_ALIGN_PARAGRAPH.LEFT

    # Contact
    contact = doc.add_paragraph()
    contact.add_run(f"Email: {candidate['email']}\n").font.size = Pt(10)
    contact.add_run(f"Phone: {candidate['phone']}\n").font.size = Pt(10)
    contact.add_run(f"Location: {candidate['location']}").font.size = Pt(10)

    # Summary
    doc.add_heading("Summary", level=1)
    doc.add_paragraph(candidate["summary"])

    # Experience
    doc.add_heading("Experience", level=1)
    for line in candidate["experience"]:
        if line.startswith("-"):
            doc.add_paragraph(line, style="List Bullet")
        else:
            p = doc.add_paragraph()
            p.add_run(line).bold = True

    # Education
    doc.add_heading("Education", level=1)
    doc.add_paragraph(candidate["education"])

    # Skills
    doc.add_heading("Skills", level=1)
    doc.add_paragraph(candidate["skills"])

    doc.save(str(output_path))


def main() -> None:
    print(f"Generating {len(CANDIDATES)} test resumes...\n")
    for i, candidate in enumerate(CANDIDATES, start=1):
        filename = f"{i:02d}_{candidate['name'].replace(' ', '_')}.docx"
        path = OUTPUT_DIR / filename
        create_resume_docx(candidate, path)
        print(f"  [{candidate['match']:>14}] {filename}")

    print(f"\n✅ Done. {len(CANDIDATES)} files in {OUTPUT_DIR.absolute()}")
    print("\nNext steps:")
    print("1. Login to http://localhost:3000 as B2B user")
    print("2. Go to /b2b/batches")
    print("3. Click 'New Batch'")
    print(f"4. Upload files from: {OUTPUT_DIR.absolute()}")


if __name__ == "__main__":
    main()