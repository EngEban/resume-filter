# ============================================================
# app/services/excel_export.py
# Generate Excel workbooks from batch results.
# ============================================================
import io
from datetime import datetime, timezone
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

# ---------- Styling ----------
HEADER_BG = "2563EB"
HEADER_FG = "FFFFFF"
ALT_ROW_BG = "F1F5F9"
BORDER_COLOR = "CBD5E1"

LEVEL_COLORS = {
    "excellent": "15803D",
    "good": "16A34A",
    "average": "F59E0B",
    "below_average": "EA580C",
    "weak": "DC2626",
}

COLUMNS: list[tuple[str, int]] = [
    ("Rank", 6),
    ("Candidate Name", 28),
    ("Email", 30),
    ("Phone", 18),
    ("ATS Score", 11),
    ("Level", 14),
    ("Matched Keywords", 40),
    ("Missing Keywords", 40),
    ("Status", 12),
    ("File Name", 32),
    ("Uploaded At", 20),
]


def _level_from_score(score: float) -> str:
    if score >= 90:
        return "excellent"
    if score >= 75:
        return "good"
    if score >= 60:
        return "average"
    if score >= 40:
        return "below_average"
    return "weak"


def _thin_border() -> Border:
    side = Side(style="thin", color=BORDER_COLOR)
    return Border(left=side, right=side, top=side, bottom=side)


def _format_keywords(values: Any) -> str:
    if not values:
        return ""
    if isinstance(values, list):
        return ", ".join(str(v) for v in values[:25])
    return str(values)


def build_batch_workbook(batch: dict, resumes: list[dict]) -> bytes:
    """
    Build an Excel workbook for a batch.

    Args:
        batch: dict with id, job_title, status, etc.
        resumes: list of resume dicts from the API.

    Returns:
        Raw .xlsx file bytes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Candidates"

    # ---------- Title row ----------
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    title_cell = ws.cell(row=1, column=1)
    title_cell.value = f"Batch: {batch.get('job_title', 'Untitled')}"
    title_cell.font = Font(size=14, bold=True, color="0F172A")
    title_cell.alignment = Alignment(horizontal="left", vertical="center")
    ws.row_dimensions[1].height = 24

    # ---------- Metadata row ----------
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(COLUMNS))
    meta_cell = ws.cell(row=2, column=1)
    meta_cell.value = (
        f"Batch ID: {batch.get('id', '—')}  |  "
        f"Status: {batch.get('status', '—')}  |  "
        f"Total: {batch.get('total_resumes', 0)}  |  "
        f"Completed: {batch.get('completed', 0)}  |  "
        f"Failed: {batch.get('failed', 0)}  |  "
        f"Exported: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M UTC')}"
    )
    meta_cell.font = Font(size=10, color="64748B")
    ws.row_dimensions[2].height = 18

    # ---------- Header row ----------
    header_row = 4
    for idx, (name, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=idx, value=name)
        cell.font = Font(bold=True, color=HEADER_FG, size=11)
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )
        cell.border = _thin_border()
        ws.column_dimensions[get_column_letter(idx)].width = width
    ws.row_dimensions[header_row].height = 26

    # ---------- Sort resumes ----------
    sorted_resumes = sorted(
        resumes,
        key=lambda r: (r.get("ats_score") is None, -(r.get("ats_score") or 0)),
    )

    # ---------- Data rows ----------
    border = _thin_border()
    for rank, resume in enumerate(sorted_resumes, start=1):
        row_idx = header_row + rank
        score = resume.get("ats_score")
        level = _level_from_score(float(score)) if score is not None else "weak"

        values = [
            rank,
            resume.get("candidate_name") or "—",
            resume.get("candidate_email") or "—",
            resume.get("candidate_phone") or "—",
            round(float(score), 1) if score is not None else "—",
            level.replace("_", " ").title() if score is not None else "—",
            _format_keywords(resume.get("matched_keywords")),
            _format_keywords(resume.get("missing_keywords")),
            resume.get("status") or "—",
            resume.get("file_name") or "—",
            (resume.get("created_at") or "")[:19].replace("T", " "),
        ]

        for col_idx, value in enumerate(values, start=1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.border = border
            cell.alignment = Alignment(
                vertical="center",
                wrap_text=col_idx in (7, 8),
                horizontal="center" if col_idx in (1, 5, 6, 9, 11) else "left",
            )

            if rank % 2 == 0:
                cell.fill = PatternFill("solid", fgColor=ALT_ROW_BG)

            if col_idx == 5 and score is not None:
                cell.font = Font(bold=True, color=LEVEL_COLORS[level])

        ws.row_dimensions[row_idx].height = 20

    # ---------- Freeze + autofilter ----------
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    ws.auto_filter.ref = (
        f"A{header_row}:"
        f"{get_column_letter(len(COLUMNS))}{header_row + len(sorted_resumes)}"
    )

    # ---------- Bytes ----------
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ============================================================
# Comparison workbook (multi-candidate side-by-side)
# ============================================================
def build_comparison_workbook(
    batch: dict,
    resumes: list[dict],
) -> bytes:
    """
    Build a comparison workbook where each candidate is a column.

    Useful for quickly comparing top candidates.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison"

    top = sorted(
        resumes,
        key=lambda r: (r.get("ats_score") is None, -(r.get("ats_score") or 0)),
    )[:10]

    if not top:
        ws.cell(row=1, column=1, value="No candidates to compare.")
        buf = io.BytesIO()
        wb.save(buf)
        return buf.getvalue()

    # ---------- Header row: candidate names ----------
    ws.cell(row=1, column=1, value="Metric").font = Font(bold=True, color=HEADER_FG)
    ws.cell(row=1, column=1).fill = PatternFill("solid", fgColor=HEADER_BG)
    ws.cell(row=1, column=1).alignment = Alignment(horizontal="center")
    ws.column_dimensions["A"].width = 26

    for idx, r in enumerate(top, start=2):
        cell = ws.cell(
            row=1,
            column=idx,
            value=(r.get("candidate_name") or "—")[:25],
        )
        cell.font = Font(bold=True, color=HEADER_FG)
        cell.fill = PatternFill("solid", fgColor=HEADER_BG)
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(idx)].width = 24

    rows = [
        ("ATS Score", lambda r: round(float(r["ats_score"]), 1) if r.get("ats_score") is not None else "—"),
        ("Level", lambda r: _level_from_score(float(r["ats_score"])) if r.get("ats_score") is not None else "—"),
        ("Email", lambda r: r.get("candidate_email") or "—"),
        ("Phone", lambda r: r.get("candidate_phone") or "—"),
        ("Status", lambda r: r.get("status") or "—"),
        ("Missing Keywords", lambda r: _format_keywords(r.get("missing_keywords"))),
    ]

    border = _thin_border()
    for row_idx, (label, extractor) in enumerate(rows, start=2):
        label_cell = ws.cell(row=row_idx, column=1, value=label)
        label_cell.font = Font(bold=True)
        label_cell.fill = PatternFill("solid", fgColor=ALT_ROW_BG)
        label_cell.border = border

        for col_idx, r in enumerate(top, start=2):
            cell = ws.cell(row=row_idx, column=col_idx, value=extractor(r))
            cell.border = border
            cell.alignment = Alignment(vertical="center", wrap_text=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()