from io import BytesIO

from openpyxl import load_workbook

from app.services.excel_export import build_batch_workbook, build_comparison_workbook
from app.services.pdf_export import build_batch_pdf, build_resume_pdf


def _sample_resume() -> dict:
    return {
        "candidate_name": "Jane Doe",
        "candidate_email": "jane@example.com",
        "candidate_phone": "+970599123456",
        "ats_score": 87.5,
        "matched_keywords": ["python", "fastapi"],
        "missing_keywords": ["kubernetes"],
        "status": "completed",
        "file_name": "jane-doe.pdf",
        "created_at": "2026-10-10T12:00:00+00:00",
        "score_breakdown": {"keyword_match": 82.0},
        "suggestions": [
            {
                "type": "add_keyword",
                "priority": "high",
                "description": "Add Kubernetes experience when applicable.",
            }
        ],
    }


def test_batch_workbook_is_valid_xlsx():
    content = build_batch_workbook(
        {"id": "batch-1", "job_title": "Backend Engineer", "status": "completed"},
        [_sample_resume()],
    )

    assert content.startswith(b"PK")
    workbook = load_workbook(BytesIO(content), read_only=True)
    assert workbook.sheetnames == ["Candidates"]
    assert workbook["Candidates"]["A1"].value == "Batch: Backend Engineer"


def test_comparison_workbook_is_valid_xlsx():
    content = build_comparison_workbook(
        {"id": "batch-1", "job_title": "Backend Engineer", "status": "completed"},
        [_sample_resume()],
    )

    assert content.startswith(b"PK")
    workbook = load_workbook(BytesIO(content), read_only=True)
    assert workbook.sheetnames == ["Comparison"]
    assert workbook["Comparison"]["A1"].value == "Metric"


def test_batch_pdf_has_pdf_signature():
    content = build_batch_pdf(
        {"id": "batch-1", "job_title": "Backend Engineer", "status": "completed"},
        [_sample_resume()],
    )

    assert content.startswith(b"%PDF-")
    assert b"%%EOF" in content


def test_resume_pdf_has_pdf_signature():
    content = build_resume_pdf(_sample_resume())

    assert content.startswith(b"%PDF-")
    assert b"%%EOF" in content
